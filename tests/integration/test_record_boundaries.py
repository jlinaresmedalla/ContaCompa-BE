"""Transaction, correction ownership and export regressions."""

import hashlib
import io
from decimal import Decimal
from pathlib import Path
from uuid import UUID

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from contacompa.application.services.intake import submit
from contacompa.application.services.purchase_docs import materialize
from contacompa.config import Settings
from contacompa.domain.config import RunConfig
from contacompa.domain.purchase_doc import SourceKind
from contacompa.entrypoints.api.app import create_app
from contacompa.infrastructure.blob import LocalBlobStore, sha256_hex
from contacompa.infrastructure.db import queue
from contacompa.infrastructure.db.models import (
    Company,
    Document,
    ExtractionResult,
    Job,
    JobStatus,
    PurchaseDoc,
    Supplier,
)
from tests.conftest import make_invoice_pdf
from tests.integration.conftest import make_company

pytestmark = pytest.mark.integration


async def _record(session: AsyncSession, company: Company, number: str, name: str) -> UUID:
    key = hashlib.sha256(number.encode()).hexdigest()
    document = Document(
        company_id=company.id,
        sha256=key,
        filename=f"{number}.pdf",
        mime_type="application/pdf",
        source_kind=SourceKind.PDF_TEXT,
        size_bytes=1,
        storage_key=key,
    )
    session.add(document)
    await session.flush()
    job = Job(document_id=document.id, config={}, status=JobStatus.DONE)
    session.add(job)
    await session.flush()
    fields = {
        "supplier_ruc": {"value": "20600000005"},
        "supplier_name": {"value": name},
        "doc_number": {"value": number},
        "issue_date": {"value": "2026-09-01"},
        "currency": {"value": "PEN"},
        "total_amount": {"value": "100.00"},
        "prices_include_igv": {"value": "true"},
        "buyer_ruc": {"value": company.ruc},
        "lines": [
            {
                "description": {"value": "Paper"},
                "quantity": {"value": "2"},
                "unit": {"value": "unit"},
                "unit_price": {"value": "50.00"},
                "line_total": {"value": "100.00"},
            }
        ],
    }
    result = ExtractionResult(
        job_id=job.id,
        document_id=document.id,
        schema_type="purchase_doc",
        schema_version="1",
        fields=fields,
        missing=[],
        provider="synthetic",
        model="none",
        prompt_version="test",
        parser="bundled",
        input_tokens=0,
        output_tokens=0,
        cost_usd=Decimal("0"),
        latency_ms=0,
        turnaround_ms=0,
    )
    session.add(result)
    await session.flush()
    record = await materialize(
        session,
        company=company,
        document=document,
        result_id=result.id,
        doc_type="invoice",
        fields=fields,
    )
    await session.commit()
    return record.id


async def test_failed_enqueue_rolls_back_document_and_cleans_blob(
    sessions: async_sessionmaker[AsyncSession],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings = Settings(
        _env_file=None,
        database_url="postgresql+asyncpg://unused/unused",
        api_key="test-key",
        blob_dir=str(tmp_path),
    )
    data = make_invoice_pdf()
    key = sha256_hex(data)
    blobs = LocalBlobStore(tmp_path)

    async def fail_enqueue(*_: object, **__: object) -> Job:
        raise RuntimeError("enqueue failed")

    monkeypatch.setattr(queue, "enqueue", fail_enqueue)
    async with sessions() as session:
        company = await make_company(session)
        with pytest.raises(RuntimeError, match="enqueue failed"):
            await submit(
                session,
                blobs,
                settings,
                company,
                data=data,
                filename="invoice.pdf",
                run_config=RunConfig(
                    provider="gemini", model="test", parser="native", prompt_version="v1"
                ),
            )
        assert await session.scalar(select(Document.id).where(Document.sha256 == key)) is None
    assert not await blobs.exists(key)


async def test_supplier_name_correction_isolated_conflict_and_export(
    sessions: async_sessionmaker[AsyncSession], pg_url: str, tmp_path: Path
) -> None:
    settings = Settings(
        _env_file=None, database_url=pg_url, api_key="test-key", blob_dir=str(tmp_path)
    )
    app = create_app(settings, telemetry=False)
    async with app.router.lifespan_context(app):
        async with sessions() as session:
            company = await make_company(session)
            first = await _record(session, company, "F001-000001", "First printed name")
            second = await _record(session, company, "F001-000002", "Second printed name")
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            headers = {"X-API-Key": "test-key"}
            changed = await client.patch(
                f"/v1/purchase-docs/{first}",
                headers=headers,
                json={"fields": {"supplier_name": "Corrected on first"}},
            )
            assert changed.status_code == 200, changed.text
            assert changed.json()["supplier"]["legal_name"] == "Corrected on first"
            other = await client.get(f"/v1/purchase-docs/{second}", headers=headers)
            assert other.json()["supplier"]["legal_name"] == "Second printed name"
            conflict = await client.patch(
                f"/v1/purchase-docs/{second}",
                headers=headers,
                json={"fields": {"doc_number": "F001-000001"}},
            )
            assert conflict.status_code == 409
            after = await client.get(f"/v1/purchase-docs/{second}", headers=headers)
            assert after.json()["doc_number"] == "F001-000002"
            assert after.json()["corrections"] == []
            exported = await client.get("/v1/exports/purchase-docs.xlsx", headers=headers)
            assert exported.status_code == 200
            book = load_workbook(io.BytesIO(exported.content), read_only=True)
            values = [cell.value for sheet in book for row in sheet for cell in row]
            assert "Corrected on first" in values
            assert "Second printed name" in values
            assert "First printed name" not in values
        async with sessions() as session:
            supplier = await session.scalar(select(Supplier).where(Supplier.ruc == "20600000005"))
            assert supplier is not None and supplier.legal_name == "First printed name"
            assert (
                await session.scalar(
                    select(PurchaseDoc.supplier_name).where(PurchaseDoc.id == second)
                )
                == "Second printed name"
            )
