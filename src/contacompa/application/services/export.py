"""Excel report of purchase docs.

Sheet 1 is a master-detail report: one band per purchase doc (supplier, number, base, IGV, total,
observations) with its lines underneath, grouped so Excel can fold them. Sheet 2 is the same lines
flat, one row each, for filters and pivot tables. Prices appear with and without IGV (ADR 0008)."""

import io
from collections.abc import Sequence
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Literal
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import Outline
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from contacompa.application.services.records import lines_by_doc
from contacompa.domain.igv import CENTS, UNIT_PRICE_PLACES, both_prices, doc_totals
from contacompa.infrastructure.db.models import PurchaseDoc, PurchaseDocLine, Supplier

Language = Literal["es", "en"]

LABELS: dict[Language, dict[str, str]] = {
    "es": {
        "title": "Comprobantes de compra",
        "period": "Periodo",
        "all_dates": "todas las fechas",
        "generated": "Generado",
        "count": "comprobantes",
        "report": "Reporte",
        "flat": "Items",
        "number": "Número",
        "date": "Fecha",
        "supplier": "Proveedor",
        "ruc": "RUC",
        "type": "Tipo",
        "currency": "Moneda",
        "taxable": "Base (sin IGV)",
        "igv": "IGV",
        "total": "Total (con IGV)",
        "observations": "Observaciones",
        "line": "#",
        "description": "Descripción",
        "quantity": "Cant.",
        "unit": "Unidad",
        "price_without": "P. unit. sin IGV",
        "price_with": "P. unit. con IGV",
        "total_without": "Importe sin IGV",
        "total_with": "Importe con IGV",
        "printed": "Precios impresos",
        "printed_true": "con IGV",
        "printed_false": "sin IGV",
        "printed_none": "no se sabe",
        "grand_total": "Total general",
        "invoice": "Factura",
        "sales_receipt": "Boleta",
        "sales_note": "Nota de venta",
        "credit_note": "Nota de crédito",
        "other": "Otro",
    },
    "en": {
        "title": "Purchase docs",
        "period": "Period",
        "all_dates": "all dates",
        "generated": "Generated",
        "count": "purchase docs",
        "report": "Report",
        "flat": "Lines",
        "number": "Number",
        "date": "Date",
        "supplier": "Supplier",
        "ruc": "RUC",
        "type": "Type",
        "currency": "Currency",
        "taxable": "Base (ex IGV)",
        "igv": "IGV",
        "total": "Total (inc IGV)",
        "observations": "Observations",
        "line": "#",
        "description": "Description",
        "quantity": "Qty",
        "unit": "Unit",
        "price_without": "Unit price ex IGV",
        "price_with": "Unit price inc IGV",
        "total_without": "Line total ex IGV",
        "total_with": "Line total inc IGV",
        "printed": "Printed prices",
        "printed_true": "inc IGV",
        "printed_false": "ex IGV",
        "printed_none": "unknown",
        "grand_total": "Grand total",
        "invoice": "Invoice",
        "sales_receipt": "Sales receipt",
        "sales_note": "Sales note",
        "credit_note": "Credit note",
        "other": "Other",
    },
}

NAVY = "1F3A5F"
BAND = PatternFill("solid", fgColor="DCE6F1")
HEADER = PatternFill("solid", fgColor=NAVY)
TOTAL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="B7C4D6")
MONEY = "#,##0.00"
PRICE = "#,##0.0000"
QTY = "#,##0.####"
DATE = "DD/MM/YYYY"
WIDTHS = (16, 12, 44, 14, 16, 16, 16, 16, 16, 50)

Row = tuple[PurchaseDoc, Supplier | None]
CellValue = str | int | Decimal | date | None


async def export_xlsx(
    session: AsyncSession,
    company_id: UUID,
    company_name: str,
    date_from: date | None,
    date_to: date | None,
    language: Language = "es",
) -> tuple[bytes, int]:
    """Workbook bytes and count; every exported record is marked as exported."""
    stmt = (
        select(PurchaseDoc, Supplier)
        .outerjoin(Supplier, Supplier.id == PurchaseDoc.supplier_id)
        .where(PurchaseDoc.company_id == company_id)
        .order_by(PurchaseDoc.issue_date, PurchaseDoc.doc_number)
    )
    if date_from is not None:
        stmt = stmt.where(PurchaseDoc.issue_date >= date_from)
    if date_to is not None:
        stmt = stmt.where(PurchaseDoc.issue_date <= date_to)
    rows: list[Row] = [(row[0], row[1]) for row in (await session.execute(stmt)).all()]
    lines = await lines_by_doc(session, [doc.id for doc, _ in rows])

    labels = LABELS[language]
    workbook = Workbook()
    report = workbook.active
    assert report is not None
    report.title = labels["report"]
    _report_sheet(report, rows, lines, labels, company_name, date_from, date_to)
    _flat_sheet(workbook.create_sheet(labels["flat"]), rows, lines, labels)
    buffer = io.BytesIO()
    workbook.save(buffer)

    now = datetime.now(UTC)
    for doc, _ in rows:
        doc.exported_at = now
    await session.commit()
    return buffer.getvalue(), len(rows)


def _report_sheet(
    sheet: Worksheet,
    rows: Sequence[Row],
    lines: dict[UUID, list[PurchaseDocLine]],
    labels: dict[str, str],
    company_name: str,
    date_from: date | None,
    date_to: date | None,
) -> None:
    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_properties.outlinePr = Outline(summaryBelow=False)  # fold button on the band
    sheet.sheet_view.showGridLines = False

    period = f"{_fmt(date_from)} - {_fmt(date_to)}" if date_from or date_to else labels["all_dates"]
    sheet["A1"] = f"{labels['title']} · {company_name}"
    sheet["A1"].font = Font(size=16, bold=True, color=NAVY)
    sheet["A2"] = (
        f"{labels['period']}: {period}   ·   {labels['generated']}: "
        f"{datetime.now(UTC):%d/%m/%Y %H:%M} UTC   ·   {len(rows)} {labels['count']}"
    )
    sheet["A2"].font = Font(size=10, color="595959")

    header = [
        labels[key]
        for key in (
            "number",
            "date",
            "supplier",
            "ruc",
            "type",
            "currency",
            "taxable",
            "igv",
            "total",
            "observations",
        )
    ]
    _header_row(sheet, 4, header)
    sheet.freeze_panes = "A5"

    row = 5
    grand: dict[str, list[Decimal]] = {}
    for doc, supplier in rows:
        totals = doc_totals(doc.total_amount)
        band: list[CellValue] = [
            doc.doc_number,
            doc.issue_date,
            doc.supplier_name,
            supplier.ruc if supplier else None,
            labels.get(doc.doc_type.value, doc.doc_type.value),
            doc.currency,
            totals.taxable,
            totals.igv,
            totals.total,
            _observations(doc),
        ]
        for column, value in enumerate(band, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.fill = BAND
            cell.font = Font(bold=True, color=NAVY)
            cell.border = Border(top=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=column == 10)
        sheet.cell(row=row, column=2).number_format = DATE
        for column in (7, 8, 9):
            sheet.cell(row=row, column=column).number_format = MONEY
        if totals.total is not None and doc.currency:
            sums = grand.setdefault(doc.currency, [Decimal(0)] * 3)
            for i, value in enumerate((totals.taxable, totals.igv, totals.total)):
                sums[i] += value or Decimal(0)
        row += 1

        sub_header = [
            "",
            labels["line"],
            labels["description"],
            labels["quantity"],
            labels["unit"],
            labels["price_without"],
            labels["price_with"],
            labels["total_without"],
            labels["total_with"],
            f"{labels['printed']}: {_printed(doc.prices_include_igv, labels)}",
        ]
        for column, value in enumerate(sub_header, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = Font(size=9, italic=True, color="7F7F7F")
            cell.alignment = Alignment(horizontal="left" if column in (3, 10) else "right")
        sheet.row_dimensions[row].outline_level = 1
        row += 1

        for line in lines.get(doc.id, []):
            price = both_prices(line.unit_price, doc.prices_include_igv, UNIT_PRICE_PLACES)
            amount = both_prices(line.line_total, doc.prices_include_igv, CENTS)
            line_values: list[CellValue] = [
                None,
                line.line_number,
                line.description,
                line.quantity,
                line.unit,
                price.without_igv,
                price.with_igv,
                amount.without_igv,
                amount.with_igv,
                None,
            ]
            if doc.prices_include_igv is None:  # nothing to derive: show what was printed
                line_values[5], line_values[7] = line.unit_price, line.line_total
            for column, line_value in enumerate(line_values, start=1):
                cell = sheet.cell(row=row, column=column, value=line_value)
                cell.font = Font(size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=column == 3)
            sheet.cell(row=row, column=4).number_format = QTY
            for column in (6, 7):
                sheet.cell(row=row, column=column).number_format = PRICE
            for column in (8, 9):
                sheet.cell(row=row, column=column).number_format = MONEY
            sheet.row_dimensions[row].outline_level = 1
            row += 1
        row += 1  # breathing space between purchase docs

    for currency, sums in sorted(grand.items()):
        total_values: list[CellValue] = [labels["grand_total"], None, None, None, None, currency]
        for column, total_value in enumerate([*total_values, *sums], start=1):
            cell = sheet.cell(row=row, column=column, value=total_value)
            cell.fill = TOTAL
            cell.font = Font(bold=True)
            cell.border = Border(top=Side(style="medium", color=NAVY))
        for column in (7, 8, 9):
            sheet.cell(row=row, column=column).number_format = MONEY
        row += 1


def _flat_sheet(
    sheet: Worksheet,
    rows: Sequence[Row],
    lines: dict[UUID, list[PurchaseDocLine]],
    labels: dict[str, str],
) -> None:
    keys = (
        "number",
        "date",
        "ruc",
        "supplier",
        "type",
        "currency",
        "line",
        "description",
        "quantity",
        "unit",
        "printed",
        "price_without",
        "price_with",
        "total_without",
        "total_with",
        "observations",
    )
    widths = (16, 12, 14, 40, 14, 10, 5, 44, 10, 10, 14, 16, 16, 16, 16, 50)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    _header_row(sheet, 1, [labels[key] for key in keys])
    sheet.freeze_panes = "A2"
    row = 2
    for doc, supplier in rows:
        for line in lines.get(doc.id, []):
            price = both_prices(line.unit_price, doc.prices_include_igv, UNIT_PRICE_PLACES)
            amount = both_prices(line.line_total, doc.prices_include_igv, CENTS)
            sheet.append(
                [
                    doc.doc_number,
                    doc.issue_date,
                    supplier.ruc if supplier else None,
                    doc.supplier_name,
                    labels.get(doc.doc_type.value, doc.doc_type.value),
                    doc.currency,
                    line.line_number,
                    line.description,
                    line.quantity,
                    line.unit,
                    _printed(doc.prices_include_igv, labels),
                    price.without_igv,
                    price.with_igv,
                    amount.without_igv,
                    amount.with_igv,
                    _observations(doc),
                ]
            )
            sheet.cell(row=row, column=2).number_format = DATE
            sheet.cell(row=row, column=9).number_format = QTY
            for column in (12, 13):
                sheet.cell(row=row, column=column).number_format = PRICE
            for column in (14, 15):
                sheet.cell(row=row, column=column).number_format = MONEY
            row += 1
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{max(row - 1, 1)}"


def _header_row(sheet: Worksheet, row: int, values: list[str]) -> None:
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.fill = HEADER
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 30


def _observations(doc: PurchaseDoc) -> str | None:
    return (
        "\n".join(f"• {issue.get('code')}: {issue.get('detail')}" for issue in doc.issues) or None
    )


def _printed(includes_igv: bool | None, labels: dict[str, str]) -> str:
    return labels[
        {True: "printed_true", False: "printed_false", None: "printed_none"}[includes_igv]
    ]


def _fmt(value: date | None) -> str:
    return value.strftime("%d/%m/%Y") if value else "…"
